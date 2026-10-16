import csv
from collections.abc import Callable
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db.models import Q
from django.http import HttpResponse
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import extend_schema
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.exceptions import NotFound
from rest_framework.generics import ListAPIView
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.accounts.permissions import IsAdmin, IsStaffOrAdmin
from apps.accounts.roles import BusinessRole
from apps.exchanges.models import ExchangeRequest
from apps.operations.models import AuditLog
from apps.operations.selectors import (
    get_dashboard,
    get_inventory_report,
    get_recommendation_report,
    get_returns_report,
    get_sales_report,
)
from apps.operations.services import create_staff, update_staff
from common.serializers import ApiErrorSerializer

from .serializers import (
    AuditLogFilterSerializer,
    AuditLogOutputSerializer,
    CreateStaffInputSerializer,
    ReportOutputSerializer,
    ReportRangeInputSerializer,
    StaffOutputSerializer,
    UpdateStaffInputSerializer,
)

User = get_user_model()


def _report_params(request) -> dict:
    serializer = ReportRangeInputSerializer(data=request.query_params)
    serializer.is_valid(raise_exception=True)
    return serializer.validated_data


def _csv_safe(value) -> str:
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def _csv_response(*, filename: str, headers: list[str], rows: list[list]) -> HttpResponse:
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_csv_safe(value) for value in row])
    return response


def _add_workbook_sheet(workbook, *, title: str, headers: list[str], rows: list[list]):
    sheet = workbook.create_sheet(title=title)
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="0F766E")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, column in enumerate(sheet.columns, start=1):
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
    return sheet


class DashboardView(APIView):
    permission_classes = [IsStaffOrAdmin]

    @extend_schema(parameters=[ReportRangeInputSerializer], responses={200: ReportOutputSerializer})
    def get(self, request):
        params = _report_params(request)
        include_finance = (
            request.user.groups.filter(name=BusinessRole.ADMIN).exists()
            or request.user.is_superuser
        )
        return Response(
            get_dashboard(
                start=params["start"],
                end=params["end"],
                timezone_name=params["timezone"],
                include_finance=include_finance,
            )
        )


class BaseReportView(APIView):
    permission_classes = [IsStaffOrAdmin]
    selector: Callable

    @extend_schema(parameters=[ReportRangeInputSerializer], responses={200: ReportOutputSerializer})
    def get(self, request):
        params = _report_params(request)
        kwargs = {"start": params["start"], "end": params["end"]}
        if self.selector in (get_sales_report,):
            kwargs["timezone_name"] = params["timezone"]
        return Response(self.selector(**kwargs))


class SalesReportView(BaseReportView):
    permission_classes = [IsAdmin]
    selector = staticmethod(get_sales_report)


class InventoryReportView(BaseReportView):
    selector = staticmethod(get_inventory_report)


class ReturnsReportView(BaseReportView):
    selector = staticmethod(get_returns_report)


class RecommendationReportView(BaseReportView):
    selector = staticmethod(get_recommendation_report)


class SalesExportView(APIView):
    permission_classes = [IsAdmin]

    @extend_schema(
        parameters=[ReportRangeInputSerializer],
        responses={(200, "text/csv"): OpenApiTypes.BINARY},
    )
    def get(self, request):
        params = _report_params(request)
        report = get_sales_report(
            start=params["start"], end=params["end"], timezone_name=params["timezone"]
        )
        rows = [[row["date"], row["orders"], row["revenue"]] for row in report["daily"]]
        return _csv_response(
            filename="bao-cao-doanh-thu.csv", headers=["Ngày", "Số đơn", "Doanh thu VND"], rows=rows
        )


class InventoryExportView(APIView):
    permission_classes = [IsStaffOrAdmin]

    @extend_schema(
        parameters=[ReportRangeInputSerializer],
        responses={(200, "text/csv"): OpenApiTypes.BINARY},
    )
    def get(self, request):
        params = _report_params(request)
        report = get_inventory_report(start=params["start"], end=params["end"])
        rows = [
            [row["sku"], row["product_name"], row["quantity"], row["threshold"]]
            for row in report["low_stock"]
        ]
        return _csv_response(
            filename="bao-cao-ton-kho.csv", headers=["SKU", "Sản phẩm", "Tồn", "Ngưỡng"], rows=rows
        )


class ReturnsExportView(APIView):
    permission_classes = [IsStaffOrAdmin]

    @extend_schema(
        parameters=[ReportRangeInputSerializer],
        responses={(200, "text/csv"): OpenApiTypes.BINARY},
    )
    def get(self, request):
        params = _report_params(request)
        report = get_returns_report(start=params["start"], end=params["end"])
        rows = [[row["status"], row["count"]] for row in report["by_status"]]
        return _csv_response(
            filename="bao-cao-doi-hang.csv", headers=["Trạng thái", "Số yêu cầu"], rows=rows
        )


class RecommendationExportView(APIView):
    permission_classes = [IsStaffOrAdmin]

    @extend_schema(
        parameters=[ReportRangeInputSerializer],
        responses={(200, "text/csv"): OpenApiTypes.BINARY},
    )
    def get(self, request):
        params = _report_params(request)
        report = get_recommendation_report(start=params["start"], end=params["end"])
        rows = [
            [
                row[key]
                for key in (
                    "strategy",
                    "requests",
                    "impressions",
                    "clicks",
                    "ctr_percent",
                    "add_to_carts",
                    "purchases",
                    "coverage_percent",
                )
            ]
            for row in report["strategies"]
        ]
        return _csv_response(
            filename="bao-cao-goi-y.csv",
            headers=[
                "Chiến lược",
                "Request",
                "Impression",
                "Click",
                "CTR %",
                "Thêm giỏ",
                "Mua",
                "Coverage %",
            ],
            rows=rows,
        )


class ConsolidatedExcelExportView(APIView):
    permission_classes = [IsAdmin]

    @extend_schema(
        parameters=[ReportRangeInputSerializer],
        responses={
            (
                200,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ): OpenApiTypes.BINARY
        },
    )
    def get(self, request):
        params = _report_params(request)
        sales = get_sales_report(
            start=params["start"], end=params["end"], timezone_name=params["timezone"]
        )
        inventory = get_inventory_report(start=params["start"], end=params["end"])
        returns = get_returns_report(start=params["start"], end=params["end"])
        recommendations = get_recommendation_report(start=params["start"], end=params["end"])
        workbook = Workbook()
        workbook.remove(workbook.active)
        summary = sales["summary"]
        summary_sheet = _add_workbook_sheet(
            workbook,
            title="Tổng quan",
            headers=["Chỉ số", "Giá trị"],
            rows=[
                ["Doanh thu thuần", int(Decimal(summary["net_revenue"]))],
                ["Giá vốn", int(Decimal(summary["cost_of_goods_sold"]))],
                ["Lợi nhuận", int(Decimal(summary["profit"]))],
                ["Đơn đã thanh toán", summary["paid_orders"]],
                ["Đơn đã hủy", summary["cancelled_orders"]],
            ],
        )
        summary_sheet["B2"].number_format = '#,##0 "₫"'
        summary_sheet["B3"].number_format = '#,##0 "₫"'
        summary_sheet["B4"].number_format = '#,##0 "₫"'
        sales_sheet = _add_workbook_sheet(
            workbook,
            title="Doanh thu theo ngày",
            headers=["Ngày", "Số đơn", "Doanh thu"],
            rows=[
                [date.fromisoformat(row["date"]), row["orders"], int(Decimal(row["revenue"]))]
                for row in sales["daily"]
            ],
        )
        for cell in sales_sheet["C"][1:]:
            cell.number_format = '#,##0 "₫"'
        for cell in sales_sheet["A"][1:]:
            cell.number_format = "dd/mm/yyyy"
        best_sellers_sheet = _add_workbook_sheet(
            workbook,
            title="Sản phẩm bán chạy",
            headers=["Mã sản phẩm", "Tên sản phẩm", "Số lượng", "Doanh thu thuần"],
            rows=[
                [
                    row["product_id"],
                    row["product_name"],
                    row["quantity"],
                    int(Decimal(row["net_revenue"])),
                ]
                for row in sales["best_sellers"]
            ],
        )
        for cell in best_sellers_sheet["D"][1:]:
            cell.number_format = '#,##0 "₫"'
        _add_workbook_sheet(
            workbook,
            title="Tồn kho thấp",
            headers=["Mã SKU", "Sản phẩm", "Tồn kho", "Ngưỡng cảnh báo"],
            rows=[
                [row["sku"], row["product_name"], row["quantity"], row["threshold"]]
                for row in inventory["low_stock"]
            ],
        )
        _add_workbook_sheet(
            workbook,
            title="Đổi hàng",
            headers=["Trạng thái", "Số yêu cầu"],
            rows=[
                [ExchangeRequest.Status(row["status"]).label, row["count"]]
                for row in returns["by_status"]
            ],
        )
        _add_workbook_sheet(
            workbook,
            title="Gợi ý sản phẩm",
            headers=["Chiến lược", "Lượt gợi ý", "Hiển thị", "Nhấp", "CTR (%)", "Mua hàng"],
            rows=[
                [
                    {
                        "popular": "Phổ biến",
                        "personalized": "Cá nhân hóa",
                        "similar": "Sản phẩm tương tự",
                        "frequently_bought_together": "Thường mua cùng",
                    }.get(row["strategy"], row["strategy"]),
                    row["requests"],
                    row["impressions"],
                    row["clicks"],
                    Decimal(row["ctr_percent"]),
                    row["purchases"],
                ]
                for row in recommendations["strategies"]
            ],
        )
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="bao-cao-tong-hop.xlsx"'
        return response


class StaffListCreateView(APIView):
    permission_classes = [IsAdmin]

    @extend_schema(responses={200: StaffOutputSerializer(many=True)})
    def get(self, request):
        staff = (
            User.objects.filter(groups__name=BusinessRole.STAFF).distinct().order_by("-date_joined")
        )
        if search := request.query_params.get("search", "").strip():
            staff = staff.filter(Q(email__icontains=search) | Q(full_name__icontains=search))
        if request.query_params.get("is_active") in {"true", "false"}:
            staff = staff.filter(is_active=request.query_params["is_active"] == "true")
        paginator = PageNumberPagination()
        page = paginator.paginate_queryset(staff, request, view=self)
        return paginator.get_paginated_response(StaffOutputSerializer(page, many=True).data)

    @extend_schema(
        request=CreateStaffInputSerializer,
        responses={201: StaffOutputSerializer, 400: ApiErrorSerializer},
    )
    def post(self, request):
        serializer = CreateStaffInputSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        staff = create_staff(**serializer.validated_data)
        return Response(StaffOutputSerializer(staff).data, status=status.HTTP_201_CREATED)


class StaffDetailView(APIView):
    permission_classes = [IsAdmin]

    @extend_schema(
        request=UpdateStaffInputSerializer,
        responses={200: StaffOutputSerializer, 400: ApiErrorSerializer, 404: ApiErrorSerializer},
    )
    def patch(self, request, staff_id: int):
        staff = User.objects.filter(pk=staff_id, groups__name=BusinessRole.STAFF).first()
        if not staff:
            raise NotFound("Không tìm thấy nhân viên.")
        serializer = UpdateStaffInputSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        staff = update_staff(staff=staff, **serializer.validated_data)
        return Response(StaffOutputSerializer(staff).data)


class AuditLogListView(ListAPIView):
    permission_classes = [IsAdmin]
    serializer_class = AuditLogOutputSerializer

    def get_queryset(self):
        filters = AuditLogFilterSerializer(data=self.request.query_params)
        filters.is_valid(raise_exception=True)
        queryset = AuditLog.objects.select_related("actor")
        data = filters.validated_data
        if action := data.get("action"):
            queryset = queryset.filter(action=action)
        if entity_type := data.get("entity_type"):
            queryset = queryset.filter(entity_type=entity_type)
        if actor_email := data.get("actor_email"):
            queryset = queryset.filter(actor__email__iexact=actor_email)
        if start_date := data.get("start_date"):
            queryset = queryset.filter(created_at__date__gte=start_date)
        if end_date := data.get("end_date"):
            queryset = queryset.filter(created_at__date__lte=end_date)
        if search := data.get("search", "").strip():
            queryset = queryset.filter(
                Q(summary__icontains=search)
                | Q(actor__email__icontains=search)
                | Q(entity_id__icontains=search)
            )
        return queryset
